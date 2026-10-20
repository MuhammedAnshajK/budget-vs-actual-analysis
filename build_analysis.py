"""Builds the Budget vs Actual Analysis workbook.

Reads data/budget.csv and data/actual.csv and writes budget_vs_actual_2025.xlsx with:
monthly Budget and Actual P&Ls, a full-year variance summary with materiality flags,
price / volume / cost analysis, an operating profit bridge (with waterfall chart),
monthly variance trends, written commentary on the major deviations, and checks.
Every analysis cell is an Excel formula, so changing an input updates everything.

Run:  python generate_data.py && python build_analysis.py
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
OUT = ROOT / "budget_vs_actual_2025.xlsx"
MONTHS = [f"2025-{m:02d}" for m in range(1, 13)]
MLABEL = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MCOLS = [get_column_letter(2 + i) for i in range(12)]   # B..M
FY = "N"
FIRST_M, LAST_M = MCOLS[0], MCOLS[-1]
PRODUCTS = ["Running Shoes", "Apparel", "Accessories"]
OPEX = ["Marketing", "Salaries and benefits", "Rent and occupancy", "Freight and logistics",
        "IT and software", "Travel", "Other admin"]

ARIAL = "Arial"
F = Font(name=ARIAL, size=10)
F_IN = Font(name=ARIAL, size=10, color="0000FF")
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_B = Font(name=ARIAL, size=10, bold=True)
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=14, bold=True, color="1F4E79")
F_SEC = Font(name=ARIAL, size=10, bold=True, color="1F4E79")
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", fgColor="1F4E79")
FILL_SEC = PatternFill("solid", fgColor="DCE6F1")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_FAV = PatternFill("solid", fgColor="E2EFDA")
FILL_ADV = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="808080")
TOP = Border(top=THIN)
TOTAL = Border(top=THIN, bottom=Side(style="double", color="808080"))

K = '#,##0;(#,##0);"-"'
K1 = '#,##0.0;(#,##0.0);"-"'
USD = '$#,##0.00'
PCT = '0.0%;(0.0%);"-"'
PTS = '+0.0" pts";-0.0" pts";"-"'


def load(name):
    out = {}
    with open(ROOT / "data" / f"{name}.csv", newline="") as f:
        for row in csv.DictReader(f):
            out.setdefault(row["line"], {})[row["month"]] = float(row["value"])
    return out


def title(ws, text, sub):
    ws["A1"], ws["A2"] = text, sub
    ws["A1"].font, ws["A2"].font = F_TITLE, F_NOTE
    ws.sheet_view.showGridLines = False


def head_row(ws, r, labels, start=1, left=(0,)):
    for i, t in enumerate(labels):
        c = ws.cell(row=r, column=start + i, value=t)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="left" if i in left else "right", wrap_text=True, vertical="bottom")


def section(ws, r, text, width):
    for c in range(1, width + 1):
        ws.cell(row=r, column=c).fill = FILL_SEC
    ws.cell(row=r, column=1, value=text).font = F_SEC


# ---------- Budget / Actual sheets (identical layout) ----------
ROW = {}


def plan_rows():
    r = 5
    ROW["sec_vol"] = r; r += 1
    for p in PRODUCTS:
        for k in ("units", "price", "unit cost"):
            ROW[(p, k)] = r; r += 1
    r += 1
    ROW["sec_pl"] = r; r += 1
    for p in PRODUCTS:
        ROW[("rev", p)] = r; r += 1
    ROW["rev"] = r; r += 1
    for p in PRODUCTS:
        ROW[("cogs", p)] = r; r += 1
    ROW["cogs"] = r; r += 1
    ROW["gp"] = r; r += 1
    ROW["gm"] = r; r += 2
    for o in OPEX:
        ROW[("opex", o)] = r; r += 1
    ROW["opex"] = r; r += 1
    ROW["op"] = r; r += 1
    ROW["opm"] = r


def write_plan(ws, data, label):
    title(ws, f"{label}: monthly P&L, 2025",
          "Units in thousands, prices and unit costs in $, all other amounts in $ thousands. Blue = input.")
    head_row(ws, 4, ["$ thousands"] + MLABEL + ["FY 2025"])
    section(ws, ROW["sec_vol"], "Volumes, prices and unit costs", 14)
    for p in PRODUCTS:
        for k, fmt, lab in (("units", K1, "units (k)"), ("price", USD, "selling price ($)"), ("unit cost", USD, "unit cost ($)")):
            r = ROW[(p, k)]
            ws.cell(row=r, column=1, value=f"{p} {lab}").font = F
            for i, col in enumerate(MCOLS):
                c = ws[f"{col}{r}"]
                c.value, c.font, c.number_format = data[f"{p} {k}"][MONTHS[i]], F_IN, fmt
            fy = ws[f"{FY}{r}"]
            if k == "units":
                fy.value = f"=SUM({FIRST_M}{r}:{LAST_M}{r})"
            elif k == "price":   # weighted average: revenue / units
                fy.value = f"={FY}{ROW[('rev', p)]}/{FY}{ROW[(p, 'units')]}"
            else:
                fy.value = f"={FY}{ROW[('cogs', p)]}/{FY}{ROW[(p, 'units')]}"
            fy.number_format, fy.font = fmt, F_B
    section(ws, ROW["sec_pl"], "Income statement", 14)
    lines = [(("rev", p), f"Revenue: {p}", "rev") for p in PRODUCTS] + [("rev", "Total revenue", "sum_rev")] + \
            [(("cogs", p), f"Cost of sales: {p}", "cogs") for p in PRODUCTS] + [("cogs", "Total cost of sales", "sum_cogs")] + \
            [("gp", "Gross profit", "gp"), ("gm", "Gross margin %", "gm")] + \
            [(("opex", o), o, "opex") for o in OPEX] + [("opex", "Total operating expenses", "sum_opex"),
                                                        ("op", "Operating profit", "op"), ("opm", "Operating margin %", "opm")]
    for key, lab, kind in lines:
        r = ROW[key]
        bold = kind.startswith("sum") or kind in ("gp", "op")
        ws.cell(row=r, column=1, value=lab).font = F_B if bold else F
        for i, col in enumerate(MCOLS + [FY]):
            c = ws[f"{col}{r}"]
            if kind == "rev":
                p = key[1]
                c.value = f"={col}{ROW[(p, 'units')]}*{col}{ROW[(p, 'price')]}" if col != FY else f"=SUM({FIRST_M}{r}:{LAST_M}{r})"
            elif kind == "cogs":
                p = key[1]
                c.value = f"={col}{ROW[(p, 'units')]}*{col}{ROW[(p, 'unit cost')]}" if col != FY else f"=SUM({FIRST_M}{r}:{LAST_M}{r})"
            elif kind == "sum_rev":
                c.value = f"=SUM({col}{ROW[('rev', PRODUCTS[0])]}:{col}{ROW[('rev', PRODUCTS[-1])]})"
            elif kind == "sum_cogs":
                c.value = f"=SUM({col}{ROW[('cogs', PRODUCTS[0])]}:{col}{ROW[('cogs', PRODUCTS[-1])]})"
            elif kind == "gp":
                c.value = f"={col}{ROW['rev']}-{col}{ROW['cogs']}"
            elif kind == "gm":
                c.value = f"={col}{ROW['gp']}/{col}{ROW['rev']}"
            elif kind == "opex":
                if col == FY:
                    c.value = f"=SUM({FIRST_M}{r}:{LAST_M}{r})"
                else:
                    c.value, c.font = data[key[1]][MONTHS[i]], F_IN
            elif kind == "sum_opex":
                c.value = f"=SUM({col}{ROW[('opex', OPEX[0])]}:{col}{ROW[('opex', OPEX[-1])]})"
            elif kind == "op":
                c.value = f"={col}{ROW['gp']}-{col}{ROW['opex']}"
            elif kind == "opm":
                c.value = f"={col}{ROW['op']}/{col}{ROW['rev']}"
            if c.font != F_IN:
                c.font = F_B if (bold or col == FY) else F
            c.number_format = PCT if kind in ("gm", "opm") else K
            if kind.startswith("sum") or kind == "gp":
                c.border = TOP
            if kind == "op":
                c.border = TOTAL
    ws.column_dimensions["A"].width = 34
    for col in MCOLS:
        ws.column_dimensions[col].width = 9
    ws.column_dimensions[FY].width = 11
    ws.freeze_panes = "B5"


def B(key, col=FY):
    return f"Budget!{col}{ROW[key]}"


def A(key, col=FY):
    return f"Actual!{col}{ROW[key]}"


def rng(sheet, key):
    return f"{sheet}!{FIRST_M}{ROW[key]}:{LAST_M}{ROW[key]}"


def build():
    budget, actual = load("budget"), load("actual")
    plan_rows()
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    ws_s = wb.create_sheet("Summary")
    ws_c = wb.create_sheet("Commentary")
    ws_r = wb.create_sheet("Revenue Analysis")
    ws_x = wb.create_sheet("Cost Analysis")
    ws_m = wb.create_sheet("Monthly")
    ws_b = wb.create_sheet("Budget")
    ws_a = wb.create_sheet("Actual")
    ws_k = wb.create_sheet("Checks")
    write_plan(ws_b, budget, "Budget")
    write_plan(ws_a, actual, "Actual")

    # ---------- Summary ----------
    title(ws_s, "Budget vs Actual: Full Year 2025",
          "Illustrative sportswear company, $ thousands. Variance is shown so that positive = favourable (better than budget) and negative = unfavourable.")
    ws_s["A4"], ws_s["A5"] = "Materiality: flag if |variance %| is at least", "and |variance| is at least ($k)"
    ws_s["A4"].font = ws_s["A5"].font = F
    ws_s["E4"], ws_s["E5"] = 0.05, 100
    for c, fmt in ((ws_s["E4"], '0%'), (ws_s["E5"], K)):
        c.font, c.fill, c.number_format = F_IN, FILL_KEY, fmt
    THR_P, THR_K = "Summary!$E$4", "Summary!$E$5"
    head_row(ws_s, 7, ["$ thousands", "Budget", "Actual", "Variance F/(U)", "Variance %", "Status"])
    lines = [(("rev", p), f"Revenue: {p}", 1) for p in PRODUCTS] + [("rev", "Total revenue", 1)] + \
            [(("cogs", p), f"Cost of sales: {p}", -1) for p in PRODUCTS] + [("cogs", "Total cost of sales", -1),
                                                                            ("gp", "Gross profit", 1), ("gm", "Gross margin %", 0)] + \
            [(("opex", o), o, -1) for o in OPEX] + [("opex", "Total operating expenses", -1),
                                                    ("op", "Operating profit", 1), ("opm", "Operating margin %", 0)]
    SROW = {}
    r = 8
    flag_first = r
    for key, lab, sign in lines:
        bold = key in ("rev", "cogs", "gp", "opex", "op")
        ws_s.cell(row=r, column=1, value=lab).font = F_B if bold else F
        ws_s[f"B{r}"], ws_s[f"C{r}"] = "=" + B(key), "=" + A(key)
        if sign == 0:
            ws_s[f"D{r}"] = f"=(C{r}-B{r})*100"
            ws_s[f"D{r}"].number_format = PTS
            ws_s[f"B{r}"].number_format = ws_s[f"C{r}"].number_format = PCT
        else:
            ws_s[f"D{r}"] = f"=C{r}-B{r}" if sign > 0 else f"=B{r}-C{r}"
            ws_s[f"E{r}"] = f"=IF(B{r}=0,0,D{r}/ABS(B{r}))"
            ws_s[f"F{r}"] = (f'=IF(AND(ABS(E{r})>={THR_P},ABS(D{r})>={THR_K}),'
                             f'IF(D{r}>=0,"Favourable: review","Adverse: investigate"),IF(ROUND(D{r},0)=0,"On budget",IF(D{r}>0,"Favourable","Unfavourable")))')
            for col in "BCD":
                ws_s[f"{col}{r}"].number_format = K
            ws_s[f"E{r}"].number_format = PCT
        for col in "BC":
            ws_s[f"{col}{r}"].font = F_LINK
        for col in "DEF":
            ws_s[f"{col}{r}"].font = F_B if bold else F
        if bold:
            for col in "ABCDEF":
                ws_s[f"{col}{r}"].border = TOP
        SROW[key] = r
        r += 1
    flag_last = r - 1
    ws_s.conditional_formatting.add(f"F{flag_first}:F{flag_last}",
                                    FormulaRule(formula=[f'LEFT(F{flag_first},7)="Adverse"'], fill=FILL_ADV))
    ws_s.conditional_formatting.add(f"F{flag_first}:F{flag_last}",
                                    FormulaRule(formula=[f'F{flag_first}="Favourable: review"'], fill=FILL_FAV))

    # Operating profit bridge
    r += 1
    section(ws_s, r, "Operating profit bridge: budget to actual", 6)
    r += 1
    head_row(ws_s, r, ["Step", "Impact F/(U)", "Running total", "Base", "Increase", "Decrease"])
    r += 1
    bridge_first = r
    vol = "+".join(f"SUMPRODUCT({rng('Actual', (p, 'units'))}-{rng('Budget', (p, 'units'))},{rng('Budget', (p, 'price'))}-{rng('Budget', (p, 'unit cost'))})" for p in PRODUCTS)
    price = "+".join(f"SUMPRODUCT({rng('Actual', (p, 'price'))}-{rng('Budget', (p, 'price'))},{rng('Actual', (p, 'units'))})" for p in PRODUCTS)
    cost = "+".join(f"SUMPRODUCT({rng('Budget', (p, 'unit cost'))}-{rng('Actual', (p, 'unit cost'))},{rng('Actual', (p, 'units'))})" for p in PRODUCTS)
    steps = [("Budget operating profit", "=" + B("op"), "total"),
             ("Sales volume (at budget margin)", "=" + vol, "step"),
             ("Selling prices", "=" + price, "step"),
             ("Product unit costs", "=" + cost, "step")] + \
            [(o, f"={B(('opex', o))}-{A(('opex', o))}", "step") for o in OPEX] + \
            [("Actual operating profit", "=" + A("op"), "total")]
    for i, (lab, f, kind) in enumerate(steps):
        ws_s.cell(row=r, column=1, value=lab).font = F_B if kind == "total" else F
        ws_s[f"B{r}"] = f
        ws_s[f"B{r}"].font = F_LINK if kind == "total" else F
        if kind == "total":
            ws_s[f"C{r}"] = f"=B{r}"
            ws_s[f"D{r}"] = 0
            ws_s[f"E{r}"] = f"=B{r}"     # totals drawn as full bars in the Increase column, recoloured below
            ws_s[f"F{r}"] = 0
        else:
            ws_s[f"C{r}"] = f"=C{r-1}+B{r}"
            ws_s[f"D{r}"] = f"=MIN(C{r-1},C{r})"
            ws_s[f"E{r}"] = f"=MAX(B{r},0)"
            ws_s[f"F{r}"] = f"=MAX(-B{r},0)"
        for col in "BCDEF":
            ws_s[f"{col}{r}"].number_format = K
        for col in "DEF":
            ws_s[f"{col}{r}"].font = F_NOTE
        r += 1
    bridge_last = r - 1
    ws_s.cell(row=r, column=1, value="Bridge check (should be 0)").font = F_NOTE
    ws_s[f"B{r}"] = f"=ROUND(C{bridge_last-1}-B{bridge_last},2)"
    ws_s[f"B{r}"].number_format = K
    bridge_check = f"Summary!B{r}"
    ws_s.cell(row=bridge_first - 1, column=4).value = "Chart helper: base"
    r += 2

    # waterfall chart (stacked bars with an invisible base)
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
    for col, color in (("D", None), ("E", "2A78D6"), ("F", "E34948")):
        ref_ = Reference(ws_s, min_col=ord(col) - 64, min_row=bridge_first - 1, max_row=bridge_last)
        ch.add_data(ref_, titles_from_data=True)
        s = ch.series[-1]
        if color is None:
            s.graphicalProperties.noFill = True
            s.graphicalProperties.line.noFill = True
        else:
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color
    from openpyxl.chart.marker import DataPoint
    for idx in (0, bridge_last - bridge_first):   # budget and actual totals in grey
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = "898781"
        pt.graphicalProperties.line.solidFill = "898781"
        ch.series[1].dPt.append(pt)
    ch.set_categories(Reference(ws_s, min_col=1, min_row=bridge_first, max_row=bridge_last))
    ch.title = "Operating profit bridge, budget to actual ($k)"
    ch.legend = None
    ch.y_axis.numFmt = "#,##0"
    ch.y_axis.scaling.min = 10000
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E1E0D9"))
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.gapWidth = 40
    ch.height, ch.width = 9, 24
    ws_s.add_chart(ch, f"H{bridge_first - 2}")

    ws_s.column_dimensions["A"].width = 34
    for col in "BCDE":
        ws_s.column_dimensions[col].width = 13
    ws_s.column_dimensions["F"].width = 20
    ws_s.freeze_panes = "B8"

    # ---------- Revenue analysis ----------
    title(ws_r, "Revenue Analysis: price and volume",
          "Revenue variance split into volume (selling more or fewer units, at budget price) and price (selling at a different price). Positive = favourable.")
    head_row(ws_r, 4, ["Product", "Budget units (k)", "Actual units (k)", "Volume %", "Budget avg price", "Actual avg price", "Price %",
                       "Budget revenue", "Actual revenue", "Volume variance", "Price variance", "Total variance", "Check"])
    r = 5
    for p in PRODUCTS:
        ws_r.cell(row=r, column=1, value=p).font = F
        ws_r[f"B{r}"], ws_r[f"C{r}"] = "=" + B((p, "units")), "=" + A((p, "units"))
        ws_r[f"D{r}"] = f"=C{r}/B{r}-1"
        ws_r[f"E{r}"], ws_r[f"F{r}"] = "=" + B((p, "price")), "=" + A((p, "price"))
        ws_r[f"G{r}"] = f"=F{r}/E{r}-1"
        ws_r[f"H{r}"], ws_r[f"I{r}"] = "=" + B(("rev", p)), "=" + A(("rev", p))
        ws_r[f"J{r}"] = f"=SUMPRODUCT({rng('Actual', (p, 'units'))}-{rng('Budget', (p, 'units'))},{rng('Budget', (p, 'price'))})"
        ws_r[f"K{r}"] = f"=SUMPRODUCT({rng('Actual', (p, 'price'))}-{rng('Budget', (p, 'price'))},{rng('Actual', (p, 'units'))})"
        ws_r[f"L{r}"] = f"=J{r}+K{r}"
        ws_r[f"M{r}"] = f"=ROUND(L{r}-(I{r}-H{r}),2)"
        r += 1
    ws_r.cell(row=r, column=1, value="Total").font = F_B
    for col in "BCHIJKLM":
        ws_r[f"{col}{r}"] = f"=SUM({col}5:{col}{r-1})"
        ws_r[f"{col}{r}"].font = F_B
        ws_r[f"{col}{r}"].border = TOP
    ws_r[f"D{r}"] = f"=C{r}/B{r}-1"
    rev_tot = r
    for rr in range(5, r + 1):
        for col, fmt in zip("BCDEFGHIJKLM", [K1, K1, PCT, USD, USD, PCT, K, K, K, K, K, K]):
            c = ws_r[f"{col}{rr}"]
            c.number_format = fmt
            if col in "BCEFHI" and rr < r:
                c.font = F_LINK
    ws_r[f"A{r+2}"] = "Volume variance = (actual units - budget units) x budget price.  Price variance = (actual price - budget price) x actual units. Both are calculated month by month and added up."
    ws_r[f"A{r+2}"].font = F_NOTE
    # revenue mix
    r += 4
    section(ws_r, r, "Revenue mix", 13)
    head_row(ws_r, r + 1, ["Product", "Budget mix", "Actual mix", "Change (pts)"])
    for i, p in enumerate(PRODUCTS):
        rr = r + 2 + i
        ws_r[f"A{rr}"] = p
        ws_r[f"B{rr}"] = f"=H{5+i}/H{rev_tot}"
        ws_r[f"C{rr}"] = f"=I{5+i}/I{rev_tot}"
        ws_r[f"D{rr}"] = f"=(C{rr}-B{rr})*100"
        ws_r[f"B{rr}"].number_format = ws_r[f"C{rr}"].number_format = PCT
        ws_r[f"D{rr}"].number_format = PTS
    ws_r.column_dimensions["A"].width = 20
    for i in range(2, 14):
        ws_r.column_dimensions[get_column_letter(i)].width = 12
    ws_r.row_dimensions[4].height = 30

    # ---------- Cost analysis ----------
    title(ws_x, "Cost Analysis",
          "Cost of sales split into volume and unit-cost effects, a flexible budget, and operating expenses. Positive = favourable (spent less than budget).")
    head_row(ws_x, 4, ["Cost of sales", "Budget", "Flexed budget (actual units at budget cost)", "Actual",
                       "Volume effect F/(U)", "Unit cost effect F/(U)", "Total variance F/(U)", "Check"])
    r = 5
    for p in PRODUCTS:
        ws_x.cell(row=r, column=1, value=p).font = F
        ws_x[f"B{r}"] = "=" + B(("cogs", p))
        ws_x[f"C{r}"] = f"=SUMPRODUCT({rng('Actual', (p, 'units'))},{rng('Budget', (p, 'unit cost'))})"
        ws_x[f"D{r}"] = "=" + A(("cogs", p))
        ws_x[f"E{r}"] = f"=B{r}-C{r}"
        ws_x[f"F{r}"] = f"=C{r}-D{r}"
        ws_x[f"G{r}"] = f"=E{r}+F{r}"
        ws_x[f"H{r}"] = f"=ROUND(G{r}-(B{r}-D{r}),2)"
        r += 1
    ws_x.cell(row=r, column=1, value="Total cost of sales").font = F_B
    for col in "BCDEFGH":
        ws_x[f"{col}{r}"] = f"=SUM({col}5:{col}{r-1})"
        ws_x[f"{col}{r}"].font = F_B
        ws_x[f"{col}{r}"].border = TOP
    cogs_tot = r
    for rr in range(5, r + 1):
        for col in "BCDEFGH":
            ws_x[f"{col}{rr}"].number_format = K
        if rr < r:
            ws_x[f"B{rr}"].font = ws_x[f"D{rr}"].font = F_LINK
    ws_x[f"A{r+1}"] = ("The volume effect is not a real overspend: selling more units costs more. The unit cost effect is the true cost overrun "
                       "(actual units x (actual cost - budget cost)).")
    ws_x[f"A{r+1}"].font = F_NOTE

    r += 3
    section(ws_x, r, "Operating expenses", 8)
    head_row(ws_x, r + 1, ["Expense", "Budget", "Actual", "Variance F/(U)", "Variance %", "Budget % of revenue", "Actual % of revenue", "Status"])
    r += 2
    ox_first = r
    for o in OPEX + ["Total operating expenses"]:
        key = ("opex", o) if o in OPEX else "opex"
        ws_x.cell(row=r, column=1, value=o).font = F_B if key == "opex" else F
        ws_x[f"B{r}"], ws_x[f"C{r}"] = "=" + B(key), "=" + A(key)
        ws_x[f"D{r}"] = f"=B{r}-C{r}"
        ws_x[f"E{r}"] = f"=D{r}/B{r}"
        ws_x[f"F{r}"] = f"=B{r}/{B('rev')}"
        ws_x[f"G{r}"] = f"=C{r}/{A('rev')}"
        ws_x[f"H{r}"] = f"=Summary!F{SROW[key]}"
        for col, fmt in zip("BCDEFG", [K, K, K, PCT, PCT, PCT]):
            ws_x[f"{col}{r}"].number_format = fmt
        ws_x[f"B{r}"].font = ws_x[f"C{r}"].font = F_LINK
        ws_x[f"H{r}"].font = F_LINK
        if key == "opex":
            for col in "ABCDEFGH":
                ws_x[f"{col}{r}"].border = TOP
        r += 1
    ws_x.column_dimensions["A"].width = 26
    for col in "BCDEFGH":
        ws_x.column_dimensions[col].width = 14
    ws_x.column_dimensions["H"].width = 20
    ws_x.row_dimensions[4].height = 42

    # ---------- Monthly ----------
    title(ws_m, "Monthly Variance Trend", "Monthly budget vs actual, $ thousands. Positive variance = favourable.")
    head_row(ws_m, 4, ["$ thousands"] + MLABEL + ["FY 2025"])
    mlines = [("Budget revenue", lambda c: B("rev", c), K), ("Actual revenue", lambda c: A("rev", c), K),
              ("Revenue variance F/(U)", None, K), ("Revenue variance %", None, PCT), ("", None, None),
              ("Budget operating expenses", lambda c: B("opex", c), K), ("Actual operating expenses", lambda c: A("opex", c), K),
              ("Opex variance F/(U)", None, K), ("", None, None),
              ("Budget operating profit", lambda c: B("op", c), K), ("Actual operating profit", lambda c: A("op", c), K),
              ("Operating profit variance F/(U)", None, K), ("Cumulative (YTD) profit variance", None, K),
              ("Budget gross margin %", lambda c: B("gm", c), PCT), ("Actual gross margin %", lambda c: A("gm", c), PCT)]
    MR = {}
    r = 5
    for lab, fn, fmt in mlines:
        if lab:
            ws_m.cell(row=r, column=1, value=lab).font = F_B if "variance" in lab.lower() else F
            MR[lab] = r
        r += 1
    for i, col in enumerate(MCOLS + [FY]):
        for lab, fn, fmt in mlines:
            if not lab:
                continue
            rr = MR[lab]
            c = ws_m[f"{col}{rr}"]
            if fn:
                c.value, c.font = "=" + fn(col), F_LINK
            elif lab == "Revenue variance F/(U)":
                c.value = f"={col}{MR['Actual revenue']}-{col}{MR['Budget revenue']}"
            elif lab == "Revenue variance %":
                c.value = f"={col}{MR['Revenue variance F/(U)']}/{col}{MR['Budget revenue']}"
            elif lab == "Opex variance F/(U)":
                c.value = f"={col}{MR['Budget operating expenses']}-{col}{MR['Actual operating expenses']}"
            elif lab == "Operating profit variance F/(U)":
                c.value = f"={col}{MR['Actual operating profit']}-{col}{MR['Budget operating profit']}"
            elif lab.startswith("Cumulative"):
                c.value = (f"={col}{MR['Operating profit variance F/(U)']}" if i == 0 else
                           (f"={MCOLS[i-1]}{rr}+{col}{MR['Operating profit variance F/(U)']}" if col != FY else f"={LAST_M}{rr}"))
            c.number_format = fmt
            if not fn:
                c.font = F_B if "variance" in lab.lower() else F
    ws_m.column_dimensions["A"].width = 32
    for col in MCOLS:
        ws_m.column_dimensions[col].width = 9
    ws_m.column_dimensions[FY].width = 11
    # charts
    lc = LineChart()
    for lab, color, dash in (("Budget revenue", "EB6834", "dash"), ("Actual revenue", "2A78D6", None)):
        lc.add_data(Reference(ws_m, min_col=1, max_col=13, min_row=MR[lab]), from_rows=True, titles_from_data=True)
        s = lc.series[-1]
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 28575
        if dash:
            s.graphicalProperties.line.dashStyle = dash
        s.smooth = False
    lc.set_categories(Reference(ws_m, min_col=2, max_col=13, min_row=4))
    lc.title, lc.y_axis.numFmt = "Monthly revenue: budget vs actual ($k)", "#,##0"
    lc.legend.position = "b"
    lc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E1E0D9"))
    lc.x_axis.delete = lc.y_axis.delete = False
    lc.height, lc.width = 8, 16
    ws_m.add_chart(lc, "A26")
    # helper rows so favourable and unfavourable months get their own colours
    hr = MR["Actual gross margin %"] + 2
    ws_m.cell(row=hr, column=1, value="Chart helper: favourable").font = F_NOTE
    ws_m.cell(row=hr + 1, column=1, value="Chart helper: unfavourable").font = F_NOTE
    for col in MCOLS:
        v = f"{col}{MR['Operating profit variance F/(U)']}"
        ws_m[f"{col}{hr}"], ws_m[f"{col}{hr+1}"] = f"=MAX({v},0)", f"=MIN({v},0)"
        ws_m[f"{col}{hr}"].number_format = ws_m[f"{col}{hr+1}"].number_format = K
        ws_m[f"{col}{hr}"].font = ws_m[f"{col}{hr+1}"].font = F_NOTE
    bc = BarChart()
    bc.grouping, bc.overlap = "stacked", 100
    for rr, color in ((hr, "2A78D6"), (hr + 1, "E34948")):
        bc.add_data(Reference(ws_m, min_col=1, max_col=13, min_row=rr), from_rows=True, titles_from_data=True)
        bc.series[-1].graphicalProperties.solidFill = color
        bc.series[-1].graphicalProperties.line.solidFill = color
    bc.set_categories(Reference(ws_m, min_col=2, max_col=13, min_row=4))
    bc.title, bc.y_axis.numFmt = "Monthly operating profit variance F/(U) ($k)", "#,##0"
    bc.legend = None
    bc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E1E0D9"))
    bc.x_axis.delete = bc.y_axis.delete = False
    bc.x_axis.tickLblPos = "low"
    bc.height, bc.width = 8, 16
    ws_m.add_chart(bc, "H26")

    # ---------- Commentary ----------
    title(ws_c, "Variance Commentary: the major deviations and why they happened",
          "Amounts link to the analysis sheets, so they update if the data changes. The explanations describe the events built into this illustrative data.")
    head_row(ws_c, 4, ["Area", "Variance F/(U) $k", "Variance %", "What happened", "Recommended action"], left=(0, 3, 4))
    ws_c.row_dimensions[4].height = 28
    cr = SROW
    vol_row = bridge_first + 1
    items = [
        ("Apparel sales volume", "='Revenue Analysis'!J6", "='Revenue Analysis'!D6",
         "The spring marketing campaign (March to May) lifted Apparel unit sales by about 18% from March to July, and they stayed about 8% above budget for the rest of the year. Apparel added most of the revenue upside.",
         "Keep the campaign format; build the higher Apparel run-rate into next year's budget."),
        ("Marketing", f"=Summary!D{cr[('opex', 'Marketing')]}", f"=Summary!E{cr[('opex', 'Marketing')]}",
         "The campaign cost about 40% more than budgeted marketing in March to May. It paid back: Apparel gross margin from the extra volume was well above the extra spend.",
         "Approve campaign spend through a business case so the overspend is planned rather than reported after the event."),
        ("Running Shoes unit cost", "='Cost Analysis'!F5", None,
         "The main footwear supplier raised prices by 6% from June. Selling prices were not changed, so the whole increase came out of margin.",
         "Renegotiate or dual-source; consider a price rise on the next season's range."),
        ("Freight and logistics", f"=Summary!D{cr[('opex', 'Freight and logistics')]}", f"=Summary!E{cr[('opex', 'Freight and logistics')]}",
         "Shipping rates rose in the second half, taking freight from 4.0% to about 5.5% of revenue from July.",
         "Lock in contract rates; review whether the budget's 4% of revenue assumption is still realistic."),
        ("Accessories selling price", "='Revenue Analysis'!K7", "='Revenue Analysis'!G7",
         "A 15% discount in October to December lifted Accessories volume by only about 5%, so the discount lost more revenue than it gained.",
         "Do not repeat broad discounting; use targeted promotions and measure their uplift."),
        ("IT and software", f"=Summary!D{cr[('opex', 'IT and software')]}", f"=Summary!E{cr[('opex', 'IT and software')]}",
         "A one-off $240k ERP upgrade in September was not in the budget.",
         "One-off: exclude from run-rate. Add a capital projects line to next year's budget."),
        ("Salaries and benefits", f"=Summary!D{cr[('opex', 'Salaries and benefits')]}", f"=Summary!E{cr[('opex', 'Salaries and benefits')]}",
         "Planned hires were delayed until mid-year, so salaries were below budget from January to June. This is a timing saving, not a lasting one.",
         "Treat as timing; the full cost is in the run-rate from July."),
    ]
    r = 5
    for area, v, p, what, act in items:
        ws_c.cell(row=r, column=1, value=area).font = F_B
        ws_c[f"B{r}"] = v
        ws_c[f"B{r}"].number_format, ws_c[f"B{r}"].font = K, F_LINK
        if p:
            ws_c[f"C{r}"] = p
            ws_c[f"C{r}"].number_format, ws_c[f"C{r}"].font = PCT, F_LINK
        for col, t in (("D", what), ("E", act)):
            c = ws_c[f"{col}{r}"]
            c.value, c.font, c.alignment = t, F, Alignment(wrap_text=True, vertical="top")
        for col in "ABC":
            ws_c[f"{col}{r}"].alignment = Alignment(vertical="top")
        ws_c.row_dimensions[r].height = 40
        r += 1
    r += 1
    section(ws_c, r, "Overall", 5)
    summary_text = [
        f'="Revenue was "&TEXT(Summary!D{cr["rev"]},"#,##0")&"k ("&TEXT(Summary!E{cr["rev"]},"0.0%")&") above budget, but operating profit was "&TEXT(-Summary!D{cr["op"]},"#,##0")&"k ("&TEXT(-Summary!E{cr["op"]},"0.0%")&") below budget."',
        '="The company sold more but made less: higher volumes were more than offset by higher unit costs, discounting, freight and one-off IT spend."',
        f'="Operating margin fell from "&TEXT(Summary!B{cr["opm"]},"0.0%")&" (budget) to "&TEXT(Summary!C{cr["opm"]},"0.0%")&". The IT project is one-off and the salary saving was only timing, but the supplier price rise and higher freight rates are likely to continue, so next year needs action on pricing and costs."',
    ]
    for t in summary_text:
        r += 1
        ws_c.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws_c.cell(row=r, column=1, value=t)
        c.font, c.alignment = F, Alignment(wrap_text=True, vertical="top")
        ws_c.row_dimensions[r].height = 28
    ws_c.column_dimensions["A"].width = 26
    ws_c.column_dimensions["B"].width = 14
    ws_c.column_dimensions["C"].width = 11
    ws_c.column_dimensions["D"].width = 70
    ws_c.column_dimensions["E"].width = 48

    # ---------- Checks ----------
    title(ws_k, "Checks", "Every check should be 0. Data totals are the sums of the raw CSV files.")
    head_row(ws_k, 4, ["Check", "Difference"])
    raw_rev_b = sum(budget[f"{p} units"][m] * budget[f"{p} price"][m] for p in PRODUCTS for m in MONTHS)
    raw_rev_a = sum(actual[f"{p} units"][m] * actual[f"{p} price"][m] for p in PRODUCTS for m in MONTHS)
    raw_ox_b = sum(budget[o][m] for o in OPEX for m in MONTHS)
    raw_ox_a = sum(actual[o][m] for o in OPEX for m in MONTHS)
    checks = [
        ("Budget revenue agrees to data", f"=ROUND({B('rev')}-{raw_rev_b:.4f},2)"),
        ("Actual revenue agrees to data", f"=ROUND({A('rev')}-{raw_rev_a:.4f},2)"),
        ("Budget opex agrees to data", f"=ROUND({B('opex')}-{raw_ox_b:.4f},2)"),
        ("Actual opex agrees to data", f"=ROUND({A('opex')}-{raw_ox_a:.4f},2)"),
        ("Price + volume = revenue variance", f"='Revenue Analysis'!M{rev_tot}"),
        ("Volume + unit cost = cost of sales variance", f"='Cost Analysis'!H{cogs_tot}"),
        ("Profit bridge adds up to actual profit", f"={bridge_check}"),
        ("Monthly profit variances add up to full year", f"=ROUND(Monthly!{FY}{MR['Cumulative (YTD) profit variance']}-Summary!D{cr['op']},2)"),
    ]
    for i, (lab, f) in enumerate(checks):
        ws_k.cell(row=5 + i, column=1, value=lab).font = F
        ws_k[f"B{5+i}"] = f
        ws_k[f"B{5+i}"].number_format = '#,##0.00;(#,##0.00);"-"'
    sr = 5 + len(checks) + 1
    ws_k.cell(row=sr, column=1, value="Overall status").font = F_B
    ws_k[f"B{sr}"] = f'=IF(SUMPRODUCT(ABS(B5:B{sr-2}))=0,"OK","ERROR")'
    ws_k[f"B{sr}"].font, ws_k[f"B{sr}"].fill = F_B, FILL_KEY
    ws_k.column_dimensions["A"].width = 46
    ws_k.column_dimensions["B"].width = 14

    # ---------- Cover ----------
    cover["A1"] = "Budget vs Actual Analysis, Financial Year 2025"
    cover["A1"].font = Font(name=ARIAL, size=18, bold=True, color="1F4E79")
    cover["A2"] = "Illustrative sportswear company with three product lines. $ thousands unless stated."
    cover["A2"].font = F_NOTE
    rows = [
        ("Purpose", "Compare actual 2025 results with the budget, measure revenue, expense, profit and percentage variances, and explain the major deviations."),
        ("Headline", f'="Revenue "&TEXT(Summary!E{cr["rev"]},"+0.0%;-0.0%")&" vs budget, operating profit "&TEXT(Summary!E{cr["op"]},"+0.0%;-0.0%")&": sold more, made less."'),
        ("", ""),
        ("Sheets", ""),
        ("Summary", "Full-year variance table with materiality flags, and the operating profit bridge with waterfall chart"),
        ("Commentary", "The major deviations, why they happened and what to do"),
        ("Revenue Analysis", "Revenue variance split into price and volume, and revenue mix"),
        ("Cost Analysis", "Cost of sales split into volume and unit-cost effects (flexible budget), and operating expenses"),
        ("Monthly", "Month-by-month variances, year-to-date profit variance and charts"),
        ("Budget / Actual", "Monthly P&L inputs (units, prices, unit costs, expenses)"),
        ("Checks", "Reconciliations; should say OK"),
        ("", ""),
        ("Sign convention", "Positive variance = favourable (more revenue or less cost than budget). Negative = unfavourable."),
        ("Materiality", "A line is flagged when its variance is at least 5% and at least $100k (both editable on the Summary sheet, yellow cells)."),
        ("Colour key", "Blue = input, black = formula, green = link to another sheet, yellow = key assumption"),
        ("Checks status", f"=Checks!B{sr}"),
    ]
    for i, (a, b) in enumerate(rows):
        cover.cell(row=4 + i, column=1, value=a).font = F_B
        c = cover.cell(row=4 + i, column=2, value=b)
        c.font = F
        if a == "Checks status":
            c.font, c.fill = F_B, FILL_KEY
    cover.column_dimensions["A"].width = 20
    cover.column_dimensions["B"].width = 120

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_view.showGridLines = False
    wb.save(OUT)
    print(f"Wrote {OUT.name}")


if __name__ == "__main__":
    build()
